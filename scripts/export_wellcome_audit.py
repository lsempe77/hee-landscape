"""Audit workbook of every Wellcome-funded publication, for a human check of GBD coding and theme mapping.

Run: python scripts/export_wellcome_audit.py [--raw data/raw]
Uses the same inputs and rules as build_wellcome_analysis_gbd_v5.py (it imports them), so each row matches the
page. Writes exports/wellcome_audit_<date>/wellcome_funded_audit_<date>.xlsx with one row per Wellcome-funded
publication: identifiers, title, abstract, Crossref funders, GBD 2023 coding, theme mapping and the reason a
publication is not mapped, plus empty reviewer columns. Internal use: it includes abstracts and is not
published on the site. Requires openpyxl.
"""
import argparse
import csv
from datetime import date
from pathlib import Path

import build_wellcome_analysis_gbd_v5 as w

THEME_LABELS = {k: label for k, (label, _, _) in w.PRIORITIES.items()}


def read_csv(path):
    csv.field_size_limit(10 ** 9)
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def reason(r, gbd_row, climate_decision):
    """Plain statement of how a publication was mapped, or why it is not in a theme."""
    if r["themes"]:
        parts = []
        if "infectious" in r["themes"]:
            parts.append("Infectious Diseases: GBD " + "; ".join(sorted(r["gbd"] & w.GBD_INFECTIOUS)))
        if "mental" in r["themes"]:
            parts.append("Mental Health: GBD " + "; ".join(sorted(r["gbd"] & w.GBD_MENTAL)))
        if "climate" in r["themes"]:
            parts.append(f"Climate & Health: reviewed keyword list ({r['climate_subarea']})")
        return "Mapped. " + " | ".join(parts)
    status = gbd_row.get("gbd_status", "")
    if r["gbd"]:
        why = "Coded to GBD " + "; ".join(sorted(r["gbd_l3"] or r["gbd"])) + ", which is not among the causes the themes use."
    elif status == "risk_factor_only":
        why = f"Coded to a risk factor ({gbd_row.get('gbd_risk_l3', '')}) rather than a disease, under the coding rules."
    elif status == "no_specific_disease":
        why = "Not specific to one disease, so no GBD cause was coded."
    else:
        why = f"No GBD disease code (status: {status or 'none'})."
    if climate_decision:
        why += f" Climate & Health candidate, excluded on review: {climate_decision['basis']}"
    return "Not mapped. " + why


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--raw", type=Path, default=w.ROOT / "data" / "raw")
    args = parser.parse_args()
    rows, *_ = w.load_inputs(args.raw)
    analysis = w.build_analysis(rows, {"years": [2010, 2026], "built": ""}, "")
    group_of = {s["title"]: g["label"] for g in analysis["wellcome_outside_groups"] for s in g["studies"]}
    master = {m["record_id"]: m for m in read_csv(args.raw / "hee_analysis_master_v2_1.csv") if m["is_eligible"] == "TRUE"}
    gbd = {g["record_id"]: g for g in read_csv(args.raw / w.GBD_BY_RECORD)}
    climate = {d["record_id"]: d for d in read_csv(w.CLIMATE_REVIEW)}

    columns = [("record_id", 14), ("year", 7), ("title", 60), ("abstract", 90), ("journal", 28), ("doi", 26), ("url", 30),
               ("funders_crossref_raw", 45), ("funders_identified", 40),
               ("gbd_status", 18), ("gbd_level1", 30), ("gbd_level2", 34), ("gbd_level3", 40), ("gbd_secondary_level3", 30),
               ("gbd_risk_factors", 30), ("gbd_hepatitis_recheck", 14),
               ("wellcome_themes", 26), ("mental_health_focus_condition", 14), ("climate_subarea", 26),
               ("outside_group_on_page", 30), ("mapping_reason", 70)]
    # The Audit sheet: only what a reviewer needs. Themes are mapped at GBD level 2; each study is coded at level 3.
    audit = [("ID", "record_id", 12), ("Year", "year", 7), ("Title", "title", 55), ("Abstract", "abstract", 85), ("Link", "link", 30),
             ("GBD condition coded (level 3)", "condition", 34), ("GBD group (level 2): used for theme mapping", "gbd_level2", 34),
             ("Wellcome theme", "wellcome_themes", 22), ("Why it is or isn't in a theme", "mapping_reason", 60),
             ("Is the GBD condition right?", None, 14), ("Suggested GBD condition", None, 38),
             ("Is the theme right?", None, 14), ("Suggested theme", None, 20), ("Notes", None, 40)]
    out_rows = []
    for r in sorted((r for r in rows if r["wellcome"]), key=lambda r: (r["year"], r["title"])):
        m, g = master.get(r["record_id"], {}), gbd.get(r["record_id"], {})
        cd = climate.get(r["record_id"])
        out_rows.append({
            "record_id": r["record_id"], "year": r["year"], "title": r["title"], "abstract": m.get("abstract", ""),
            "journal": m.get("journal", ""), "doi": m.get("doi", ""), "url": r["url"],
            "funders_crossref_raw": r["funders"], "funders_identified": "; ".join(r["named_funders"]),
            "gbd_status": g.get("gbd_status", ""), "gbd_level1": g.get("gbd_l1", ""), "gbd_level2": g.get("gbd_l2", ""),
            "gbd_level3": g.get("gbd_l3", ""), "gbd_secondary_level3": g.get("gbd_secondary_l3", ""),
            "gbd_risk_factors": g.get("gbd_risk_l3", ""), "gbd_hepatitis_recheck": g.get("gbd_hepatitis_recode", ""),
            "wellcome_themes": "; ".join(THEME_LABELS[k] for k in w.PRIORITIES if k in r["themes"]) or "None",
            "mental_health_focus_condition": "Yes" if r["gbd_l3"] & set(w.MENTAL_GROUPS) else "",
            "climate_subarea": r["climate_subarea"] or "",
            "outside_group_on_page": "" if r["themes"] else group_of.get(r["title"], ""),
            "mapping_reason": reason(r, g, cd if cd and cd["decision"] != "Include" else None),
            "condition": r["gbd_label"], "link": r["url"] or (f"https://doi.org/{m['doi']}" if m.get("doi") else ""),
        })

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    review_fill, header_fill = PatternFill("solid", fgColor="FFF6DC"), PatternFill("solid", fgColor="E9E2EE")
    last = len(out_rows) + 1

    ws = wb.active
    ws.title = "Audit"
    for c, (header, key, width) in enumerate(audit, 1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = Font(bold=True)
        cell.fill = review_fill if key is None else header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = width
    for i, row in enumerate(out_rows, 2):
        for c, (header, key, _) in enumerate(audit, 1):
            cell = ws.cell(row=i, column=c, value=row.get(key, "") if key else "")
            cell.alignment = Alignment(wrap_text=key in ("title", "abstract", "mapping_reason"), vertical="top")
    col = {header: ws.cell(row=1, column=c).column_letter for c, (header, _, _) in enumerate(audit, 1)}
    yes_no = DataValidation(type="list", formula1='"Yes,No,Unsure"', allow_blank=True)
    ws.add_data_validation(yes_no)
    for header in ("Is the GBD condition right?", "Is the theme right?"):
        yes_no.add(f"{col[header]}2:{col[header]}{last}")
    # Suggested condition: every GBD 2023 level-3 cause with its level-2 group (from the 'GBD 2023 causes' sheet).
    causes = sorted(read_csv(args.raw / w.GBD_HIERARCHY), key=lambda c: int(c["sort_order"]))
    level2 = sorted({c["l2_name"] for c in causes})
    dropdown = [f"{c['l3_name']} ({c['l2_name']})" for c in causes] + [f"{g} (level 2 only, no specific condition)" for g in level2]
    # The team's own groups for infections and injuries without a named disease (not in IHME's hierarchy).
    dropdown += ["Infections, disease not stated (team group; counted in Infectious Diseases)", "Injuries, cause not stated (team group)"]
    suggest = DataValidation(type="list", formula1=f"='GBD 2023 causes'!$E$2:$E${1 + len(dropdown)}", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(suggest)
    suggest.add(f"{col['Suggested GBD condition']}2:{col['Suggested GBD condition']}{last}")
    themes_dv = DataValidation(type="list", formula1='"Infectious Diseases,Mental Health,Climate & Health,None"', allow_blank=True,
                               showErrorMessage=False)
    ws.add_data_validation(themes_dv)
    themes_dv.add(f"{col['Suggested theme']}2:{col['Suggested theme']}{last}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{col['Notes']}{last}"

    details = wb.create_sheet("Full details")
    for c, (name, width) in enumerate(columns, 1):
        cell = details.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        details.column_dimensions[cell.column_letter].width = width
    for i, row in enumerate(out_rows, 2):
        for c, (name, _) in enumerate(columns, 1):
            cell = details.cell(row=i, column=c, value=row.get(name, ""))
            cell.alignment = Alignment(wrap_text=name in ("title", "abstract", "mapping_reason", "funders_crossref_raw"), vertical="top")
    details.freeze_panes = "D2"
    details.auto_filter.ref = f"A1:{details.cell(row=1, column=len(columns)).column_letter}{last}"

    guide = wb.create_sheet("How to review", 0)
    mapped = sum(1 for r in out_rows if r["wellcome_themes"] != "None")
    lines = [("Wellcome-funded publications: coding audit", True),
             (f"Built {date.today().isoformat()} by scripts/export_wellcome_audit.py from the same inputs and rules as the Wellcome page "
              f"(build_wellcome_analysis_gbd_v5.py). {len(out_rows)} publications with Wellcome Trust identified as a funder in Crossref metadata; "
              f"{mapped} mapped to at least one theme, {len(out_rows) - mapped} not mapped.", False),
             ("", False),
             ("What to check (Audit sheet)", True),
             ("1. Is the GBD condition right? Each study is coded to the specific condition (level 3) its intervention prevents or treats. "
              "If not, pick the right one under 'Suggested GBD condition'; each option shows its level-2 group, e.g. 'Cervical cancer (Neoplasms)'.", False),
             ("2. Is the theme right? Themes are mapped at GBD level 2, so the theme follows from the condition's group. "
              "If not, pick one under 'Suggested theme'.", False),
             ("Use 'Notes' for anything else. The 'GBD 2023 causes' sheet lists every condition (search with Ctrl+F); 'Coding rules' lists the rules the coders followed. "
              "'Full details' has the technical columns, for reference only.", False),
             ("", False),
             ("How themes are mapped", True),
             ("Infectious Diseases: GBD level-2 groups " + "; ".join(w.INFECTIOUS_GROUPS) + ".", False),
             ("Mental Health: GBD mental disorders or substance use disorders. mental_health_focus_condition marks Wellcome's stated focus "
              "(GBD depressive and anxiety disorders, schizophrenia, bipolar disorder), shown as a subset on the page.", False),
             ("Climate & Health: keyword screen of titles and abstracts, every candidate reviewed by hand (exports/climate_health_review_decisions_2026-10-12.csv).", False),
             ("GBD codes each study to the disease its intervention prevents or treats, not its cause: e.g. HPV vaccination to cervical cancer, "
              "hepatitis C treatment to cirrhosis and other chronic liver diseases. A risk factor is recorded on its own when no single disease is targeted.", False),
             ("", False),
             ("Columns", True),
             ("funders_crossref_raw is the funder text from Crossref; funders_identified is the cleaned list the page uses (Burroughs Wellcome Fund is a separate funder).", False),
             ("outside_group_on_page is the indicative keyword group shown in the page's pop-up for publications not mapped to a theme.", False),
             ("Internal use only: abstracts are included.", False)]
    guide.column_dimensions["A"].width = 150
    for i, (text, bold) in enumerate(lines, 1):
        cell = guide.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=13 if i == 1 else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    # GBD 2023 cause list (searchable with Ctrl+F or the filters); column E feeds the suggestion drop-down.
    cl = wb.create_sheet("GBD 2023 causes")
    cl.append(["Level 3 condition", "Level 2 group (used for theme mapping)", "Level 1", "GBD outline", "Drop-down list: condition (group)"])
    for cell in cl[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ordered = causes
    for i in range(max(len(ordered), len(dropdown))):
        c = ordered[i] if i < len(ordered) else {}
        cl.append([c.get("l3_name", ""), c.get("l2_name", ""), c.get("l1_name", ""), c.get("outline", ""), dropdown[i] if i < len(dropdown) else ""])
    for col, width in zip("ABCDE", (46, 44, 40, 12, 52)):
        cl.column_dimensions[col].width = width
    cl.freeze_panes = "A2"
    cl.auto_filter.ref = f"A1:D{1 + len(ordered)}"
    note_row = cl.max_row + 2
    cl.cell(row=note_row, column=1, value="Source: IHME GBD 2023 cause hierarchy (levels 1-3), as used by the team's coding. "
                                           "Official list and definitions: https://www.healthdata.org/research-analysis/gbd").font = Font(italic=True)

    rules = wb.create_sheet("Coding rules")
    rules.column_dimensions["A"].width = 150
    rule_lines = [
        ("GBD coding rules used for the map (team rules, 13-14 October 2026; full record in methods/gbd_coding/GBD_CODING.md)", True),
        ("1. Code the condition the evaluated intervention targets. Conditions measured only as outcomes go in the secondary field; background mentions are not coded.", False),
        ("2. Never infer a disease from a risk factor. Risk factors are recorded on their own; a cause is coded only when the text names it.", False),
        ("3. Injuries named by their nature (e.g. hip fracture) without the external cause are coded to level 1 Injuries.", False),
        ("4. Surgical and device complications are 'Adverse effects of medical treatment'. Infections are coded to the infectious disease, unless caused by human error.", False),
        ("5. General headings are coded at the level the text supports (e.g. metastasis -> level 2 Neoplasms). Conditions GBD treats as 'garbage codes' "
         "(e.g. heart failure, sepsis, infertility) get no cause (disease unattributable).", False),
        ("6. Disputed cases: give the condition its ICD-10 code and use the GBD cause IHME's ICD-10 mapping assigns to it.", False),
        ("7. A pathogen alone (e.g. 'bacterial infections') is coded to level 1 Communicable, unless it names one disease (e.g. Chlamydia -> STIs excluding HIV).", False),
        ("8. Viral hepatitis (provisional): chronic infection care -> Cirrhosis and other chronic liver diseases; prevention of new infection -> Acute hepatitis "
         "and Cirrhosis; acute infection -> Acute hepatitis.", False),
        ("", False),
        ("Examples: HPV vaccination -> Cervical cancer; hepatitis C treatment -> Cirrhosis and other chronic liver diseases; maternal Group B strep vaccination -> "
         "Neonatal disorders; alcohol pricing policy -> risk factor (alcohol use), no disease.", False),
    ]
    for i, (text, bold) in enumerate(rule_lines, 1):
        cell = rules.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=12 if i == 1 else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    out = w.ROOT / "exports" / f"wellcome_audit_{date.today().isoformat()}"
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"wellcome_funded_audit_{date.today().isoformat()}.xlsx"
    wb.save(path)
    print(f"{len(out_rows)} publications ({mapped} mapped, {len(out_rows) - mapped} not mapped) -> {path}")


if __name__ == "__main__":
    main()
